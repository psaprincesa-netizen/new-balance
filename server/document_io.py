"""Document exchange preserves supported tables, lists, links and embedded raster images."""
import base64,io,re
from PIL import Image
from markdown_it import MarkdownIt
INLINE_MARKDOWN=MarkdownIt('commonmark')
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches
from server.slides_io import image_bytes
from server.store import InvalidInput


def valid_images(images):
    if not isinstance(images,list) or len(images)>20:return False
    ids=set()
    for image in images:
        if not isinstance(image,dict) or not isinstance(image.get('id'),str) or not re.fullmatch(r'[a-zA-Z0-9_-]{1,100}',image['id']) or image['id'] in ids:return False
        ids.add(image['id'])
        if not isinstance(image.get('alt',''),str) or len(image.get('alt',''))>200:return False
        if type(image.get('width',80)) not in (int,float) or not 10<=image.get('width',80)<=100:return False
        try:image_bytes(image.get('src'))
        except InvalidInput:return False
    return True


def inline(paragraph,text):
    tokens=INLINE_MARKDOWN.parseInline(text)[0].children or []
    bold=italic=0;link=None
    for token in tokens:
        if token.type=='strong_open':bold+=1;continue
        if token.type=='strong_close':bold=max(0,bold-1);continue
        if token.type=='em_open':italic+=1;continue
        if token.type=='em_close':italic=max(0,italic-1);continue
        if token.type=='link_open':
            url=token.attrGet('href') or '';link=url if url.startswith(('http://','https://')) else None;continue
        if token.type=='link_close':link=None;continue
        value='\n' if token.type in ('softbreak','hardbreak') or (token.type=='html_inline' and re.fullmatch(r'<br\s*/?>',token.content,re.I)) else token.content
        if not value:continue
        run=paragraph.add_run(value);run.bold=bool(bold);run.italic=bool(italic)
        if token.type=='code_inline':run.font.name='Courier New'
        if link:
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),paragraph.part.relate_to(link,RT.HYPERLINK,is_external=True));h.append(run._r);paragraph._p.append(h)


def export_document(content,title,output):
    document=Document();document.core_properties.title=title;document.core_properties.author='DRAP office'
    lines=content['text'].splitlines();images={im['id']:im for im in content.get('images',[])};i=0
    def cells(line):return [x.strip().replace('\\|','|') for x in re.split(r'(?<!\\)\|',line.strip().strip('|'))]
    while i<len(lines):
        line=lines[i];i+=1
        picture=re.fullmatch(r'!\[([^\]]*)\]\(drap-image:([a-zA-Z0-9_-]+)\)',line.strip())
        if picture:
            im=images.get(picture[2])
            if not im:raise InvalidInput('Uma imagem referenciada no documento não está disponível.')
            data=image_bytes(im['src'])
            with Image.open(io.BytesIO(data)) as raster:width=min(6.5*im.get('width',80)/100,8.5*raster.width/raster.height)
            shape=document.add_picture(io.BytesIO(data),width=Inches(width));shape._inline.docPr.set('descr',im.get('alt',picture[1]));continue
        if '|' in line and i<len(lines) and re.match(r'^\s*\|?\s*:?-{3,}',lines[i]):
            headers=cells(line);i+=1;rows=[]
            while i<len(lines) and '|' in lines[i]:rows.append(cells(lines[i]));i+=1
            if len(headers)>100 or (len(rows)+1)*len(headers)>50000:raise InvalidInput('Tabela do documento excede 50 mil células ou 100 colunas.')
            table=document.add_table(rows=1,cols=len(headers));table.style='Table Grid'
            for cell,value in zip(table.rows[0].cells,headers):inline(cell.paragraphs[0],value)
            for row in rows:
                for cell,value in zip(table.add_row().cells,row):inline(cell.paragraphs[0],value)
            continue
        heading=re.match(r'^(#{1,3})\s+(.*)',line);bullet=re.match(r'^\s*([-*]|\d+\.)\s+(.*)',line)
        if heading:p=document.add_heading('',level=len(heading[1]));line=heading[2]
        elif bullet:p=document.add_paragraph(style='List Number' if bullet[1][0].isdigit() else 'List Bullet');line=bullet[2]
        else:p=document.add_paragraph()
        inline(p,line)
    document.save(output)


def import_document(stream):
    document=Document(stream);lines=[];images=[]
    def text(paragraph):
        result=[]
        for part in paragraph.iter_inner_content():
            if hasattr(part,'url'):
                result.append('['+part.text+']('+part.url.replace(')','%29')+')' if part.url.startswith(('http://','https://')) else part.text)
            else:
                value=re.sub(r'([\\`*_\[\]])',r'\\\1',part.text)
                marker='***' if part.bold and part.italic else '**' if part.bold else '*' if part.italic else ''
                result.append(marker+value+marker if value else '')
        return ''.join(result)
    for block in document.iter_inner_content():
        if hasattr(block,'rows'):
            rows=[[cell.text.replace('|','\\|').replace('\n','<br>') for cell in row.cells] for row in block.rows]
            if rows:lines.extend(['| '+' | '.join(rows[0])+' |','| '+' | '.join('---' for _ in rows[0])+' |',*['| '+' | '.join(row)+' |' for row in rows[1:]]])
            continue
        prefix=''
        if block.style:
            name=block.style.name
            if name.startswith('Heading ') and name.split()[-1].isdigit():prefix='#'*min(3,int(name.split()[-1]))+' '
            elif name.startswith('List Bullet'):prefix='- '
            elif name.startswith('List Number'):prefix='1. '
        value=text(block)
        drawings=block._p.xpath('.//w:drawing')
        if value or not drawings:lines.append(prefix+value)
        for drawing in drawings:
            blips=drawing.xpath('.//a:blip')
            for blip in blips:
                rid=blip.get(qn('r:embed'))
                if not rid:continue
                if len(images)>=20:raise InvalidInput('O documento excede 20 imagens.')
                blob=document.part.related_parts[rid].blob
                with Image.open(io.BytesIO(blob)) as im:
                    if im.width*im.height>40_000_000:raise InvalidInput('Imagem do documento excede 40 megapixels.')
                    im=im.convert('RGBA');im.thumbnail((1200,1200));out=io.BytesIO();im.save(out,format='PNG',optimize=True)
                    while len(out.getvalue())>300000 and min(im.size)>100:
                        im.thumbnail((int(im.width*.8),int(im.height*.8)));out=io.BytesIO();im.save(out,format='PNG',optimize=True)
                props=drawing.xpath('.//wp:docPr');alt=(props[0].get('descr','') if props else '')[:200]
                key='docx-'+str(len(images));image={'id':key,'src':'data:image/png;base64,'+base64.b64encode(out.getvalue()).decode(),'alt':alt,'width':80,'ratio':im.width/im.height}
                images.append(image);lines.append('!['+alt.replace(']','')+'](drap-image:'+key+')')
    if not valid_images(images):raise InvalidInput('Uma imagem do documento excede os limites.')
    result={'text':'\n'.join(lines)}
    if images:result['images']=images
    return result
