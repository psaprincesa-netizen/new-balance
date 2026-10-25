"""Raster images and editable slide objects, with bounded content and Office exchange."""
import base64
import io
import math
import re
import warnings
from PIL import Image
from server.store import InvalidInput


def image_bytes(source):
    if not isinstance(source,str) or len(source)>410000:raise InvalidInput('Imagem excede 300 KB após otimização.')
    match=re.fullmatch(r'data:image/(png|jpeg);base64,([A-Za-z0-9+/=]+)',source)
    if not match:raise InvalidInput('Use imagens PNG ou JPEG.')
    try:
        data=base64.b64decode(match[2],validate=True)
        with warnings.catch_warnings():
            warnings.simplefilter('error',Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(data)) as im:
                if im.format != {'png':'PNG','jpeg':'JPEG'}[match[1]] or im.width*im.height>4_000_000:raise ValueError()
                im.verify()
        return data
    except Exception:raise InvalidInput('Imagem inválida ou resolução excessiva.') from None


def valid_objects(objects):
    if not isinstance(objects,list) or len(objects)>30:return False
    ids=set()
    for obj in objects:
        if not isinstance(obj,dict) or not isinstance(obj.get('id'),str) or not 1<=len(obj['id'])<=100 or obj['id'] in ids:return False
        ids.add(obj['id'])
        if obj.get('type') not in ('image','text','rect','ellipse'):return False
        for k in ('x','y','w','h'):
            v=obj.get(k)
            if type(v) not in (int,float) or not math.isfinite(v) or not 0<=v<=100:return False
        if min(obj['w'],obj['h'])<1 or obj['x']+obj['w']>100.01 or obj['y']+obj['h']>100.01:return False
        if not isinstance(obj.get('alt',''),str) or len(obj.get('alt',''))>200:return False
        if obj['type']=='image':
            try:image_bytes(obj.get('src'))
            except InvalidInput:return False
        else:
            if not isinstance(obj.get('text',''),str) or len(obj.get('text',''))>10000:return False
            if not re.fullmatch(r'#[a-fA-F0-9]{6}',obj.get('color','#ffffff')) or not re.fullmatch(r'#[a-fA-F0-9]{6}',obj.get('fill','#4b5320')):return False
            if type(obj.get('fontSize',24)) not in (int,float) or not 8<=obj.get('fontSize',24)<=96:return False
            if type(obj.get('bold',False)) is not bool or obj.get('align','left') not in ('left','center','right'):return False
    return True


def export_slides(content,title,output):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.enum.text import PP_ALIGN
    p=Presentation();p.slide_width=Inches(13.333333);p.slide_height=Inches(7.5)
    p.core_properties.title=title;p.core_properties.author='DRAP office'
    theme=content.get('theme','light');color='FFFFFF' if theme!='light' else '172018'
    for entry in content['slides']:
        slide=p.slides.add_slide(p.slide_layouts[1])
        if theme!='light':
            slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('4B5320' if theme=='army' else '102237')
        for shape,text,y,h,size in [(slide.shapes.title,entry['title'],6,20,32),(slide.placeholders[1],entry['body'],30,62,22)]:
            shape.text=text;shape.left=int(p.slide_width*.06);shape.top=int(p.slide_height*y/100)
            shape.width=int(p.slide_width*(.42 if entry.get('objects') and y==30 else .88));shape.height=int(p.slide_height*h/100)
            for paragraph in shape.text_frame.paragraphs:
                paragraph.font.size=Pt(size);paragraph.font.color.rgb=RGBColor.from_string(color)
        slide.notes_slide.notes_text_frame.text=entry['notes']
        for obj in entry.get('objects',[]):
            x,y,w,h=[int(n*size/100) for n,size in zip([obj['x'],obj['y'],obj['w'],obj['h']],[p.slide_width,p.slide_height,p.slide_width,p.slide_height])]
            if obj['type']=='image':
                shape=slide.shapes.add_picture(io.BytesIO(image_bytes(obj['src'])),x,y,width=w,height=h)
                shape._element.nvPicPr.cNvPr.set('descr',obj.get('alt',''))
            else:
                if obj['type']=='text':shape=slide.shapes.add_textbox(x,y,w,h)
                else:
                    shape=slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL if obj['type']=='ellipse' else MSO_AUTO_SHAPE_TYPE.RECTANGLE,x,y,w,h)
                    shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(obj.get('fill','#4b5320')[1:]);shape.line.fill.background()
                shape.text=obj.get('text','');shape.text_frame.word_wrap=True
                for paragraph in shape.text_frame.paragraphs:
                    paragraph.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[obj.get('align','left')]
                    paragraph.font.size=Pt(obj.get('fontSize',24));paragraph.font.bold=obj.get('bold',False);paragraph.font.color.rgb=RGBColor.from_string(obj.get('color','#ffffff')[1:])
            shape.name='DRAP '+obj['type']
    p.save(output)


def import_objects(slide,presentation):
    from pptx.enum.shapes import MSO_SHAPE_TYPE,MSO_AUTO_SHAPE_TYPE
    result=[]
    for shape in slide.shapes:
        kind='image' if shape.shape_type==MSO_SHAPE_TYPE.PICTURE else shape.name[5:] if shape.name.startswith('DRAP ') else None
        if kind not in ('image','text','rect','ellipse'):continue
        x=max(0,min(99,shape.left/presentation.slide_width*100));y=max(0,min(99,shape.top/presentation.slide_height*100))
        obj={'id':'object-'+str(shape.shape_id),'type':kind,'x':x,'y':y,'w':max(1,min(100-x,shape.width/presentation.slide_width*100)),'h':max(1,min(100-y,shape.height/presentation.slide_height*100))}
        if kind=='image':
            with Image.open(io.BytesIO(shape.image.blob)) as im:
                if im.width*im.height>40_000_000:raise InvalidInput('Imagem do PPTX excede 40 megapixels.')
                im.thumbnail((1200,1200));im=im.convert('RGBA');out=io.BytesIO();im.save(out,format='PNG',optimize=True)
                while len(out.getvalue())>300000 and min(im.size)>100:
                    im.thumbnail((int(im.width*.8),int(im.height*.8)));out=io.BytesIO();im.save(out,format='PNG',optimize=True)
                obj['src']='data:image/png;base64,'+base64.b64encode(out.getvalue()).decode()
            obj['alt']=shape._element.nvPicPr.cNvPr.get('descr','')[:200]
        else:
            obj['text']=shape.text
            paragraph=shape.text_frame.paragraphs[0];font=paragraph.runs[0].font if paragraph.runs else paragraph.font
            size=font.size or paragraph.font.size
            obj['fontSize']=max(8,min(96,size.pt if size else 24));obj['bold']=bool(font.bold if font.bold is not None else paragraph.font.bold)
            obj['align']={1:'left',2:'center',3:'right'}.get(paragraph.alignment,'left')
            try:obj['color']='#'+str(font.color.rgb)
            except (AttributeError,TypeError):
                try:obj['color']='#'+str(paragraph.font.color.rgb)
                except (AttributeError,TypeError):obj['color']='#172018'
            if kind!='text':
                try:obj['fill']='#'+str(shape.fill.fore_color.rgb)
                except (AttributeError,TypeError):obj['fill']='#4b5320'
        result.append(obj)
    if not valid_objects(result):raise InvalidInput('Objetos da apresentação excedem os limites do editor.')
    return result
