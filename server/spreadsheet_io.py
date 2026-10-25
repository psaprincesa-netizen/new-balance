"""Native spreadsheet snapshots ↔ XLSX, without macros or remote resources."""
import io
from datetime import date, datetime, time
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.datetime import to_excel
from server.store import InvalidInput

BORDERS={1:'thin',2:'hair',3:'dotted',4:'dashed',5:'dashDot',6:'dashDotDot',7:'double',8:'medium',9:'mediumDashed',10:'mediumDashDot',11:'mediumDashDotDot',12:'slantDashDot',13:'thick'}

def color(value):
    if isinstance(value,dict): value=value.get('rgb')
    if isinstance(value,str):
        value=value.lstrip('#')
        if len(value) in (6,8) and all(c in '0123456789abcdefABCDEF' for c in value):return value[-6:]
    return None

def imported_color(c):
    return {'rgb':'#'+c.rgb[-6:]} if c is not None and c.type=='rgb' and isinstance(c.rgb,str) else None

def apply_style(cell, s):
    if not isinstance(s,dict):return
    cell.font=Font(name=s.get('ff','Calibri'),size=s.get('fs',11),bold=s.get('bl')==1,italic=s.get('it')==1,strike=s.get('st',{}).get('s')==1 if isinstance(s.get('st'),dict) else False,underline='single' if isinstance(s.get('ul'),dict) and s['ul'].get('s')==1 else None,color=color(s.get('cl')))
    if color(s.get('bg')):cell.fill=PatternFill('solid',fgColor=color(s['bg']))
    cell.alignment=Alignment(horizontal={1:'left',2:'center',3:'right',4:'justify'}.get(s.get('ht')),vertical={1:'top',2:'center',3:'bottom'}.get(s.get('vt')),wrap_text=s.get('tb')==3)
    if isinstance(s.get('n'),dict):cell.number_format=s['n'].get('pattern','General')
    if isinstance(s.get('bd'),dict):
        sides={}
        for k,name in [('t','top'),('b','bottom'),('l','left'),('r','right')]:
            edge=s['bd'].get(k)
            if isinstance(edge,dict):sides[name]=Side(style=BORDERS.get(edge.get('s'),'thin'),color=color(edge.get('cl')) or '000000')
        cell.border=Border(**sides)

def read_style(cell):
    f=cell.font;a=cell.alignment
    s={'ff':f.name or 'Calibri','fs':f.sz or 11,'bl':int(bool(f.b)),'it':int(bool(f.i))}
    if f.u:s['ul']={'s':1}
    if f.strike:s['st']={'s':1}
    if imported_color(f.color):s['cl']=imported_color(f.color)
    if cell.fill.patternType=='solid' and imported_color(cell.fill.fgColor):s['bg']=imported_color(cell.fill.fgColor)
    if cell.number_format!='General':s['n']={'pattern':cell.number_format}
    s.update(ht={'left':1,'center':2,'right':3,'justify':4}.get(a.horizontal,0),vt={'top':1,'center':2,'bottom':3}.get(a.vertical,0),tb=3 if a.wrap_text else 1)
    bd={}
    for k,name in [('t','top'),('b','bottom'),('l','left'),('r','right')]:
        edge=getattr(cell.border,name)
        if edge and edge.style:bd[k]={'s':next((i for i,v in BORDERS.items() if v==edge.style),1),'cl':imported_color(edge.color) or {'rgb':'#000000'}}
    if bd:s['bd']=bd
    return s

def import_workbook(stream):
    book=load_workbook(stream,read_only=False,data_only=False,keep_links=False)
    try:
        if not 1<=len(book.worksheets)<=20:raise InvalidInput('O arquivo deve ter de 1 a 20 abas.')
        snapshot={'id':'imported-workbook','name':'Importação','appVersion':'1.2.0','locale':'ptBR','dateSystem':'date1904' if book.epoch.year==1904 else 'date1900','styles':{},'sheetOrder':[],'sheets':{}}
        count=0;resources={}
        for i,ws in enumerate(book.worksheets):
            if ws.max_row>100000 or ws.max_column>1024:raise InvalidInput('A área usada excede 100 mil linhas ou 1024 colunas.')
            sid='sheet-'+str(i);snapshot['sheetOrder'].append(sid)
            data={}
            # Sparse iteration avoids allocating the entire dimension rectangle.
            for (r,c),cell in ws._cells.items():
                if cell.value is None and not cell.has_style:continue
                count+=1
                if count>50000:raise InvalidInput('Limite de 50 mil células preenchidas por arquivo.')
                value=cell.value;entry={}
                if cell.data_type=='f':entry['f']=value
                elif isinstance(value,(datetime,date,time)):
                    entry.update(v=to_excel(value,book.epoch),t=2)
                elif value is not None:entry.update(v=value,t=3 if isinstance(value,bool) else 2 if isinstance(value,(int,float)) else 1)
                if cell.has_style:entry['s']=read_style(cell)
                data.setdefault(str(r-1),{})[str(c-1)]=entry
            sheet={'id':sid,'name':ws.title,'rowCount':max(1000,ws.max_row),'columnCount':max(26,ws.max_column),'cellData':data,'defaultColumnWidth':105,'defaultRowHeight':25,'mergeData':[{'startRow':r.min_row-1,'endRow':r.max_row-1,'startColumn':r.min_col-1,'endColumn':r.max_col-1} for r in ws.merged_cells.ranges], 'rowData':{},'columnData':{}}
            for r,dim in ws.row_dimensions.items():
                if r<=100000:sheet['rowData'][str(r-1)]={'h':(dim.height or 18.75)*4/3,'hd':int(bool(dim.hidden))}
            for col,dim in ws.column_dimensions.items():
                first=dim.min or column_index_from_string(col);last=dim.max or first
                for c in range(first,min(last,1024)+1):sheet['columnData'][str(c-1)]={'w':(dim.width or 14.3)*7+5,'hd':int(bool(dim.hidden))}
            if ws.freeze_panes:
                from openpyxl.utils.cell import coordinate_to_tuple
                r,c=coordinate_to_tuple(str(ws.freeze_panes));sheet['freeze']={'xSplit':c-1,'ySplit':r-1,'startRow':r-1,'startColumn':c-1}
            snapshot['sheets'][sid]=sheet
            import_rules(ws,sid,resources)
        import json
        snapshot['resources']=[{'name':name,'data':json.dumps(data,ensure_ascii=False)} for name,data in resources.items()]
        return {'workbook':snapshot}
    finally:book.close()

def export_workbook(snapshot, output, charts=(), pivots=None):
    from server.office import excel_formula
    book=Workbook();book.remove(book.active)
    if snapshot.get('dateSystem')=='date1904':
        from openpyxl.utils.datetime import CALENDAR_MAC_1904
        book.epoch=CALENDAR_MAC_1904
    for sid in snapshot['sheetOrder']:
        data=snapshot['sheets'][sid];ws=book.create_sheet(data['name'])
        for r,cells in data.get('cellData',{}).items():
            for c,entry in cells.items():
                if not entry:continue
                cell=ws.cell(int(r)+1,int(c)+1)
                if entry.get('f'):cell.value=excel_formula(entry['f'])
                else:
                    value=entry.get('v')
                    if value is None and isinstance(entry.get('p'),dict):value=entry['p'].get('body',{}).get('dataStream','').rstrip('\r\n')
                    cell.value=value
                    if isinstance(value,str):cell.data_type='s'
                style=entry.get('s',{})
                if isinstance(style,str):style=snapshot.get('styles',{}).get(style,{})
                apply_style(cell,style)
        for merge in data.get('mergeData',[]):ws.merge_cells(start_row=merge['startRow']+1,end_row=merge['endRow']+1,start_column=merge['startColumn']+1,end_column=merge['endColumn']+1)
        for r,dim in data.get('rowData',{}).items():
            if not dim:continue
            ws.row_dimensions[int(r)+1].height=float(dim.get('h',25))*3/4
            ws.row_dimensions[int(r)+1].hidden=dim.get('hd')==1
        for c,dim in data.get('columnData',{}).items():
            if not dim:continue
            d=ws.column_dimensions[get_column_letter(int(c)+1)];d.width=max(0,(float(dim.get('w',105))-5)/7);d.hidden=dim.get('hd')==1
        freeze=data.get('freeze') or {}
        x,y=freeze.get('xSplit',0),freeze.get('ySplit',0)
        if x or y:ws.freeze_panes=ws.cell(y+1,x+1)
        ws.sheet_view.showGridLines=data.get('showGridlines',1)!=0
        export_rules(ws,sid,snapshot)
    from openpyxl.chart import BarChart,LineChart,PieChart,Reference
    from openpyxl.utils.cell import range_boundaries
    for i,definition in enumerate(charts):
        sheet=snapshot['sheets'].get(definition['sheetId'])
        if not sheet:continue
        ws=book[sheet['name']];x1,y1,x2,y2=range_boundaries(definition['range'])
        chart={'bar':BarChart,'line':LineChart,'pie':PieChart}[definition['type']]();chart.title=definition['title']
        chart.add_data(Reference(ws,min_col=x2,min_row=y1,max_row=y2),titles_from_data=True)
        chart.set_categories(Reference(ws,min_col=x1,min_row=y1+1,max_row=y2))
        chart.width=18;chart.height=10;ws.add_chart(chart,'D'+str(2+i*20))
    from server.pivots import export_pivots
    export_pivots(book,snapshot,[] if pivots is None else pivots)
    book.save(output)

def a1_range(r):
    return f"{get_column_letter(r['startColumn']+1)}{r['startRow']+1}:{get_column_letter(r['endColumn']+1)}{r['endRow']+1}"

def export_rules(ws, sid, snapshot):
    import json
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    for resource in snapshot.get('resources',[]):
        try:data=json.loads(resource.get('data') or '{}').get(sid)
        except (ValueError,AttributeError):continue
        if not data:continue
        if resource['name']=='SHEET_FILTER_PLUGIN' and isinstance(data,dict) and data.get('ref'):
            ws.auto_filter.ref=a1_range(data['ref'])
            for column in data.get('filterColumns',[]):
                values=column.get('filters')
                if values:ws.auto_filter.add_filter_column(column['colId'],values.get('filters',[]),blank=bool(values.get('blank')))
            for r in data.get('cachedFilteredOut',[]):ws.row_dimensions[r+1].hidden=True
        if resource['name']=='SHEET_DATA_VALIDATION_PLUGIN' and isinstance(data,list):
            for rule in data:
                kind=rule.get('type')
                if kind not in ('whole','decimal','list','date','time','textLength','custom'):continue
                f1=rule.get('formula1','')
                if kind=='list' and not f1.startswith('='):
                    try:values=json.loads(f1)
                    except ValueError:values=f1.split(',')
                    if not isinstance(values,list):values=f1.split(',')
                    f1='"'+','.join(str(v).replace('"','""') for v in values)+'"'
                dv=DataValidation(type=kind,operator=rule.get('operator') or None,formula1=f1,formula2=rule.get('formula2'),allow_blank=bool(rule.get('allowBlank',True)))
                dv.showErrorMessage=bool(rule.get('showErrorMessage',True));dv.errorStyle={0:'information',1:'stop',2:'warning'}.get(rule.get('errorStyle',1),'stop')
                dv.error=rule.get('error');dv.errorTitle=rule.get('errorTitle');dv.showDropDown=not rule.get('showDropDown',True)
                ws.add_data_validation(dv)
                for r in rule.get('ranges',[]):dv.add(a1_range(r))
        if resource['name']=='SHEET_CONDITIONAL_FORMATTING_PLUGIN' and isinstance(data,list):
            for entry in data:
                rule=entry.get('rule',{})
                if rule.get('type')!='highlightCell' or rule.get('subType') not in ('number','formula'):continue
                style=rule.get('style',{});fill=PatternFill('solid',fgColor=color(style.get('bg')) or 'FFFFFF') if color(style.get('bg')) else None
                font=Font(color=color(style.get('cl')),bold=style.get('bl')==1) if style else None
                if rule['subType']=='number':
                    op={'equal':'equal','notEqual':'notEqual','greaterThan':'greaterThan','greaterThanOrEqual':'greaterThanOrEqual','lessThan':'lessThan','lessThanOrEqual':'lessThanOrEqual','between':'between','notBetween':'notBetween'}.get(rule.get('operator'))
                    if not op:continue
                    values=rule.get('value',0);values=values if isinstance(values,list) else [values]
                    cf=CellIsRule(operator=op,formula=[str(v) for v in values],fill=fill,font=font,stopIfTrue=entry.get('stopIfTrue',False))
                else:cf=FormulaRule(formula=[str(rule.get('value','')).lstrip('=')],fill=fill,font=font,stopIfTrue=entry.get('stopIfTrue',False))
                for r in entry.get('ranges',[]):ws.conditional_formatting.add(a1_range(r),cf)

def import_rules(ws, sid, resources):
    import json
    def span(r):return {'startRow':r.min_row-1,'endRow':r.max_row-1,'startColumn':r.min_col-1,'endColumn':r.max_col-1}
    def add(plugin,data):resources.setdefault(plugin,{})[sid]=data
    from openpyxl.worksheet.cell_range import CellRange
    if ws.auto_filter.ref:
        columns=[]
        for col in ws.auto_filter.filterColumn:
            if col.filters:columns.append({'colId':col.colId,'filters':{'filters':list(col.filters.filter),'blank':bool(col.filters.blank)}})
        add('SHEET_FILTER_PLUGIN',{'ref':span(CellRange(ws.auto_filter.ref)),'filterColumns':columns})
    validations=[]
    for i,dv in enumerate(ws.data_validations.dataValidation):
        if dv.type not in ('whole','decimal','list','date','time','textLength','custom'):continue
        formula=str(dv.formula1 or '')
        if dv.type=='list':
            if formula.startswith('"') and formula.endswith('"'):formula=json.dumps(formula[1:-1].split(','),ensure_ascii=False)
            elif not formula.startswith('='):formula='='+formula
        validations.append({'uid':'import-validation-'+str(i),'type':dv.type,'operator':dv.operator or 'equal','formula1':formula,'formula2':str(dv.formula2 or ''),'allowBlank':bool(dv.allowBlank),'showErrorMessage':bool(dv.showErrorMessage),'errorStyle':{'information':0,'stop':1,'warning':2}.get(dv.errorStyle,1),'showDropDown':not bool(dv.showDropDown),'ranges':[span(r) for r in dv.sqref.ranges]})
    if validations:add('SHEET_DATA_VALIDATION_PLUGIN',validations)
    conditional=[]
    for region in ws.conditional_formatting:
        for rule in ws.conditional_formatting[region]:
            if rule.type not in ('cellIs','expression'):continue
            style={}
            if rule.dxf:
                if rule.dxf.fill and imported_color(rule.dxf.fill.fgColor):style['bg']=imported_color(rule.dxf.fill.fgColor)
                if rule.dxf.font:
                    if imported_color(rule.dxf.font.color):style['cl']=imported_color(rule.dxf.font.color)
                    if rule.dxf.font.bold:style['bl']=1
            config={'type':'highlightCell','style':style}
            if rule.type=='expression':config.update(subType='formula',value='='+rule.formula[0].lstrip('='))
            else:
                try:values=[float(v) for v in rule.formula]
                except ValueError:continue
                config.update(subType='number',operator=rule.operator,value=values if rule.operator in ('between','notBetween') else values[0])
            conditional.append({'cfId':'import-cf-'+str(len(conditional)),'ranges':[span(r) for r in region.sqref.ranges],'stopIfTrue':bool(rule.stopIfTrue),'rule':config})
    if conditional:add('SHEET_CONDITIONAL_FORMATTING_PLUGIN',conditional)
