"""Bounded pivot definitions and aggregation for XLSX summaries."""
import math
import re
from server.store import InvalidInput
AGGREGATES={'sum':'Soma','avg':'Média','count':'Contagem preenchida','rows':'Contagem de linhas','min':'Mínimo','max':'Máximo'}

def parse_range(address):
    m=re.fullmatch(r'([A-Z]{1,3})([1-9]\d{0,5}):([A-Z]{1,3})([1-9]\d{0,5})',address if isinstance(address,str) else '')
    if not m:raise InvalidInput('Intervalo de tabela dinâmica inválido.')
    def col(s):
        n=0
        for c in s:n=n*26+ord(c)-64
        return n-1
    a,b,c,d=m.groups();x1,x2,y1,y2=col(a),col(c),int(b)-1,int(d)-1
    height,width=y2-y1+1,x2-x1+1
    if height<2 or not 1<=width<=256 or height*width>50000 or y2>=100000 or x2>=1024:raise InvalidInput('Fonte de tabela dinâmica excede os limites.')
    return x1,y1,x2,y2

def valid_pivots(pivots):
    if not isinstance(pivots,list) or len(pivots)>10:return False
    ids=set()
    for p in pivots:
        if not isinstance(p,dict):return False
        if not isinstance(p.get('id'),str) or not 1<=len(p['id'])<=100 or p['id'] in ids:return False
        ids.add(p['id'])
        if not isinstance(p.get('title'),str) or not 1<=len(p['title'])<=100:return False
        if not isinstance(p.get('sheetId'),str) or not 1<=len(p['sheetId'])<=100:return False
        try:x1,y1,x2,y2=parse_range(p.get('range'))
        except InvalidInput:return False
        fields=[p.get('rowField'),p.get('valueField')]+[p[k] for k in ('columnField','filterField') if p.get(k) is not None]
        if not all(type(f) is int and 0<=f<x2-x1+1 for f in fields):return False
        if not isinstance(p.get('aggregate'),str) or p.get('aggregate') not in AGGREGATES or p.get('sort','source') not in ('source','asc','desc'):return False
        if not isinstance(p.get('filterValue',''),str) or len(p.get('filterValue',''))>1000:return False
    return True

def value_text(value):
    if value is None:return ''
    if type(value) is bool:return 'true' if value else 'false'
    if type(value) is float and value.is_integer():return str(int(value))
    return str(value)

def matches(value,filter_value):
    if type(value) in (int,float) and re.fullmatch(r'[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?',filter_value):
        return value==float(filter_value)
    return value_text(value)==filter_value

def calculate_pivot(source,p):
    if len(source)<2 or not source[0] or len(source)*len(source[0])>50000:raise InvalidInput('Fonte de tabela dinâmica inválida.')
    headers=[value_text(v).strip() for v in source[0]]
    if any(not v for v in headers) or len(set(headers))!=len(headers):raise InvalidInput('Cada coluna precisa de um cabeçalho único e preenchido.')
    def key(v):return ('empty','') if v is None or v=='' else ('boolean' if type(v) is bool else 'number' if type(v) in (int,float) else 'string',value_text(v))
    def label(v):return '(vazio)' if v is None or v=='' else value_text(v)
    def fresh():return {'rows':0,'count':0,'n':0,'sum':0,'min':math.inf,'max':-math.inf}
    def add(s,v):
        s['rows']+=1
        if v is not None and v!='':s['count']+=1
        if type(v) in (int,float) and math.isfinite(v):s.update(n=s['n']+1,sum=s['sum']+v,min=min(s['min'],v),max=max(s['max'],v))
    def result(s):
        op=p['aggregate']
        if op in ('rows','count','sum'):value=s[op]
        elif not s['n']:return None
        else:value=s['sum']/s['n'] if op=='avg' else s[op]
        if not math.isfinite(value):raise InvalidInput('O resultado excede o limite numérico. Reduza os valores da fonte.')
        return value
    rg={};cg={};cells={};rt={};ct={};grand=fresh()
    for row in source[1:]:
        if not any(v is not None and v!='' for v in row):continue
        if p.get('filterField') is not None and not matches(row[p['filterField']],p.get('filterValue','')):continue
        v=row[p['valueField']]
        if p['aggregate'] not in ('rows','count') and isinstance(v,str) and re.match(r'^#(?:DIV/0!|VALUE!|REF!|NAME\?|NUM!|N/A|SPILL!|CALC!)',v):raise InvalidInput('Corrija os erros de fórmula antes de exportar o resumo.')
        r=key(row[p['rowField']]);c=key(row[p['columnField']]) if p.get('columnField') is not None else ('all','')
        rg[r]=label(row[p['rowField']]);cg[c]=label(row[p['columnField']]) if p.get('columnField') is not None else AGGREGATES[p['aggregate']]
        if (len(rg)+2)*(len(cg)+2)>10000:raise InvalidInput('O resumo excede 10 mil células. Use filtros.')
        for state in (cells.setdefault((r,c),fresh()),rt.setdefault(r,fresh()),ct.setdefault(c,fresh()),grand):add(state,v)
    columns=list(cg) or [('all','')];rowkeys=list(rg)
    if p.get('sort') in ('asc','desc'):rowkeys.sort(key=lambda r:result(rt[r]) or 0,reverse=p['sort']=='desc')
    with_columns=p.get('columnField') is not None
    output=[[headers[p['rowField']],*[cg.get(c,AGGREGATES[p['aggregate']]) for c in columns],*(['Total geral'] if with_columns else [])]]
    for r in rowkeys:output.append([rg[r],*[result(cells.get((r,c),fresh())) for c in columns],*([result(rt[r])] if with_columns else [])])
    output.append(['Total geral',*[result(ct.get(c,fresh())) for c in columns],*([result(grand)] if with_columns else [])])
    return output

def pivot_source(book,p):
    sheet=book['sheets'].get(p['sheetId'])
    if not sheet:raise InvalidInput('A aba de origem da tabela dinâmica foi removida. Edite ou exclua o resumo.')
    x1,y1,x2,y2=parse_range(p['range']);data=sheet.get('cellData',{})
    if x2>=sheet['columnCount'] or y2>=sheet['rowCount']:raise InvalidInput('O intervalo da tabela dinâmica ultrapassa a aba de origem.')
    def value(r,c):
        cell=(data.get(str(r)) or {}).get(str(c)) or {}
        if cell.get('f') and cell.get('v') is None:raise InvalidInput('Recalcule e salve as fórmulas no editor antes de exportar a tabela dinâmica.')
        if 'v' in cell:return cell['v']
        rich=cell.get('p')
        body=rich.get('body') if isinstance(rich,dict) else None
        return body.get('dataStream','').rstrip('\r\n') if isinstance(body,dict) else ''
    return [[value(r,c) for c in range(x1,x2+1)] for r in range(y1,y2+1)]

def export_pivots(workbook,snapshot,pivots):
    if not valid_pivots(pivots):raise InvalidInput('Configuração de tabela dinâmica inválida.')
    from openpyxl.styles import Font,PatternFill
    for i,p in enumerate(pivots,1):
        rows=calculate_pivot(pivot_source(snapshot,p),p)
        name='Resumo '+str(i)
        suffix=1
        while name in workbook.sheetnames:
            name=f'Resumo {i} ({suffix})';suffix+=1
        sheet=workbook.create_sheet(name)
        for r,row in enumerate(rows,1):
            for c,v in enumerate(row,1):
                cell=sheet.cell(r,c,v)
                if isinstance(v,str):cell.data_type='s'
                if r in (1,len(rows)):
                    cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='4B5320')
        sheet.freeze_panes='B2'
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.print_title_rows='1:1'
