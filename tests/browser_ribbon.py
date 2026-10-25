"""Use the native spreadsheet ribbon and preserve its editing session on save."""
import io,os,tempfile,threading,uuid
from pathlib import Path
from http.server import ThreadingHTTPServer
from playwright.sync_api import sync_playwright,expect
from openpyxl import load_workbook
from server.app import handler_for
from server.store import Store


def exercise(page,base,title='Barra de ferramentas'):
    errors=[];page.on('pageerror',lambda error:errors.append(str(error)))
    page.goto(base);page.get_by_role('button',name='Criar minha conta').click()
    page.get_by_label('Nome',exact=True).fill('Ana Teste');page.get_by_label('E-mail',exact=True).fill(uuid.uuid4().hex+'@example.org');page.get_by_label('Senha',exact=True).fill(uuid.uuid4().hex+'-Teste12')
    page.get_by_role('button',name='Criar conta',exact=False).click();page.get_by_role('button',name='＋ Criar novo').click();page.locator('#create-form select').select_option('sheet');page.get_by_label('Nome do arquivo',exact=True).last.fill(title);page.get_by_role('button',name='Criar arquivo',exact=True).click()
    expect(page.locator('#spreadsheet')).to_have_attribute('data-ready','true')
    def command(name):return page.locator('[data-u-command="'+name+'"]:visible')
    def select(address):
        box=page.locator('#spreadsheet input.univer-appearance-none');box.fill(address);box.press('Enter');expect(page.locator('#sheet-address')).to_have_value(address)
    def save():
        page.get_by_role('button',name='Salvar',exact=True).click();expect(page.locator('#save-state')).to_contain_text('Versão')
    def stored():return next(x for x in page.request.get(base+'/api/items').json()['items'] if x['title']==title)
    def style(item,row,col):
        book=item['content']['workbook'];sheet=book['sheets'][book['sheetOrder'][0]];return book['styles'][sheet['cellData'][str(row)][str(col)]['s']]
    # Native name box, native typing, native ribbon: no facade or auxiliary setter.
    select('B4');page.keyboard.type('1234,5');page.keyboard.press('Enter');expect(page.locator('#sheet-address')).to_have_value('B5');select('B4');expect(page.locator('#sheet-value')).to_have_value('1234.5')
    font=command('sheet.command.set-range-font-family').locator('input');font.fill('Courier New');font.press('Enter')
    size=command('sheet.command.set-range-fontsize').locator('input');size.fill('16');size.press('Enter')
    for name in ['Negrito','Itálico','Sublinhado']:page.get_by_role('button',name=name,exact=True).click()
    page.get_by_role('button',name='Alinhamento horizontal',exact=True).click();page.get_by_role('menuitem',name='centro',exact=True).click()
    for name,index in [('sheet.command.set-range-text-color',12),('sheet.command.set-background-color',3)]:
        command(name).locator('[aria-haspopup="menu"]').click();page.locator('[role="menu"][data-state="open"] button').nth(index).click()
    command('sheet.command.set-border-basic').locator('.univerjs-icon-all-border-icon').click()
    command('sheet.command.numfmt.set.currency').click()
    select('C4');page.keyboard.type('=SOMA(1;2)');page.keyboard.press('Enter');expect(page.locator('#sheet-address')).to_have_value('C5')
    select('C5');page.keyboard.type('=SE(1=1;"SOMA; literal";"erro")');page.keyboard.press('Enter');expect(page.locator('#sheet-address')).to_have_value('C6')
    select('B4');page.get_by_role('tab',name='Dados',exact=True).click()
    # Saving used to recreate the workbook, clear undo and move back to A1/Início.
    page.locator('#spreadsheet').evaluate('el=>el.dataset.sessionMarker="kept"');save()
    expect(page.locator('#spreadsheet')).to_have_attribute('data-session-marker','kept')
    expect(page.get_by_role('tab',name='Dados',exact=True)).to_have_attribute('aria-selected','true')
    expect(page.locator('#spreadsheet input.univer-appearance-none')).to_have_value('B4')
    page.get_by_role('tab',name='Início',exact=True).click();expect(font).to_have_value('Courier New');expect(size).to_have_value('16')
    expect(page.get_by_role('button',name='Desfazer',exact=True)).to_be_enabled()
    item=stored();book=item['content']['workbook'];cells=book['sheets'][book['sheetOrder'][0]]['cellData'];s=style(item,3,1)
    assert cells['3']['1']['v']==1234.5,cells
    assert cells['3']['2']['f']=='=SUM(1,2)' and cells['3']['2']['v']==3,cells
    assert cells['4']['2']['v']=='SOMA; literal',cells
    assert s['ff']=='Courier New' and s['fs']==16 and s['bl']==1 and s['it']==1 and s['ul']['s']==1 and s['ht']==2,s
    assert s['cl'] and s['bg'] and s['n'],s
    expect(page.locator('#sheet-value')).to_have_value('1234.5')
    page.get_by_role('button',name='Desfazer',exact=True).click();page.get_by_role('button',name='Refazer',exact=True).click();save()
    # New edits must still be saved through the same mounted content reference.
    select('B4');page.get_by_role('button',name='Tachado',exact=True).click();save();assert style(stored(),3,1)['st']['s']==1
    # Clipboard menu preserves styles; "values only" must not bring formatting.
    page.context.grant_permissions(['clipboard-read','clipboard-write'],origin=base)
    select('B4');page.keyboard.press('Control+C');page.wait_for_function('async () => (await navigator.clipboard.readText()).length > 0')
    select('D4');page.get_by_role('button',name='Colar',exact=True).click();page.get_by_role('menuitem',name='Colar',exact=True).click()
    expect(page.locator('#sheet-value')).to_have_value('1234.5')
    select('E4');page.get_by_role('button',name='Colar',exact=True).click();page.get_by_role('menuitem',name='Colar valor',exact=True).click()
    page.get_by_role('tab',name='Exibir',exact=True).click();page.get_by_role('button',name='Congelar',exact=True).click();page.get_by_role('menuitem',name='Congelar primeira linha',exact=True).click()
    save();snapshot=stored()['content']['workbook'];sheet=snapshot['sheets'][snapshot['sheetOrder'][0]]
    assert sheet['cellData']['3']['3']['v']==1234.5 and style(stored(),3,3)['ff']=='Courier New'
    assert sheet['cellData']['3']['4']['v']==1234.5 and not sheet['cellData']['3']['4'].get('s')
    assert sheet['freeze']['ySplit']==1
    select('F4');page.get_by_role('tab',name='Fórmulas',exact=True).click();page.get_by_role('button',name='Funções',exact=True).click();page.get_by_role('menuitem',name='SUM',exact=True).click()
    page.keyboard.press('Enter');expect(page.locator('#sheet-address')).to_have_value('F5')
    page.get_by_role('tab',name='Início',exact=True).click()
    # Commands pushed into the overflow menu remain usable on smaller screens.
    page.set_viewport_size({'width':1280,'height':1100});select('E4')
    page.locator('[role="toolbar"] button:has(.univerjs-icon-more-vertical-icon)').click()
    command('sheet.command.numfmt.set.percent').click();page.keyboard.press('Escape')
    page.set_viewport_size({'width':1920,'height':1100});save()
    snapshot=stored()['content']['workbook'];sheet=snapshot['sheets'][snapshot['sheetOrder'][0]]
    assert sheet['cellData']['3']['5']['v']==3706.5 and '%' in style(stored(),3,4)['n']['pattern'],(sheet['cellData']['3'],style(stored(),3,4))
    # Switching cells updates the ribbon, and a second save keeps undo available.
    select('A1');expect(size).to_have_value('11');select('B4');expect(size).to_have_value('16')
    page.reload();page.get_by_role('button',name=title,exact=False).click();expect(page.locator('#spreadsheet')).to_have_attribute('data-ready','true');select('B4');expect(font).to_have_value('Courier New');expect(size).to_have_value('16')
    with page.expect_download() as dl:page.get_by_role('button',name='Baixar XLSX',exact=True).click()
    ws=load_workbook(io.BytesIO(Path(dl.value.path()).read_bytes())).worksheets[0]
    assert ws['B4'].font.name=='Courier New' and ws['B4'].font.sz==16 and ws['B4'].font.bold and ws['B4'].font.italic and ws['B4'].font.strike
    assert ws['B4'].alignment.horizontal=='center' and ws['B4'].fill.patternType=='solid' and ws['B4'].number_format!='General'
    assert ws['B4'].border.top.style and ws.freeze_panes=='A2'
    assert ws['C4'].value=='=SUM(1,2)'
    assert ws['F4'].value=='=SUM(B4:E4)' and '%' in ws['E4'].number_format
    assert not errors,errors
    page.screenshot(path='/tmp/drap-ribbon-verified.png',full_page=True)
    print('PASS: barra nativa, fonte/tamanho/cores/alinhamento/moeda, fórmulas PT, seleção/aba/desfazer preservados ao salvar, edição subsequente, reabertura e XLSX',flush=True)
    return stored()


def main():
    expect.set_options(timeout=20000)
    with tempfile.TemporaryDirectory() as tmp:
        server=ThreadingHTTPServer(('127.0.0.1',0),handler_for(Store(Path(tmp)/'test.sqlite')));threading.Thread(target=server.serve_forever,daemon=True).start()
        try:
            with sync_playwright() as p:
                browser=p.chromium.launch(headless=True,executable_path=os.environ.get('CHROMIUM_PATH','/usr/bin/chromium'));page=browser.new_page(viewport={'width':1920,'height':1100});page.set_default_timeout(20000);exercise(page,f'http://127.0.0.1:{server.server_port}');browser.close()
        finally:server.shutdown();server.server_close()
if __name__=='__main__':main()
