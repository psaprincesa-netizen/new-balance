import base64
import io
import tempfile
import unittest
import zipfile
from pathlib import Path
from server.office import import_file, export_file
from server.schema import validate_content
from server.store import InvalidInput, Store
from server.workspace import Workspace


class OfficeTest(unittest.TestCase):
    def test_real_office_roundtrips(self):
        cases = [
            ('document', 'docx', {'text': '# Plano\nConteúdo **importante**'}),
            ('sheet', 'xlsx', {'cells': [['10', '=A1*2']], 'sheets': [{'name':'Custos','cells':[['10','=A1*2']]}, {'name':'Resumo','cells':[['=Custos!B1']]}], 'activeSheet':0}),
            ('slides', 'pptx', {'slides':[{'title':'Plano','body':'Conteúdo','notes':'Lembrar'}]}),
        ]
        for kind, ext, content in cases:
            with self.subTest(kind=kind):
                body, mime = export_file({'kind':kind,'title':'Exemplo','content':content},ext)
                self.assertTrue(body.startswith(b'PK'))
                item, warnings = import_file('Exemplo.'+ext,base64.b64encode(body).decode())
                validate_content(item['kind'],item['content'])
                self.assertEqual(item['kind'],kind)
                self.assertTrue(warnings)
                if ext == 'xlsx':
                    book=item['content']['workbook']
                    self.assertEqual([book['sheets'][k]['name'] for k in book['sheetOrder']], ['Custos','Resumo'])
                    self.assertEqual(book['sheets']['sheet-0']['cellData']['0']['0']['v'],10)
                    self.assertEqual(book['sheets']['sheet-1']['cellData']['0']['0']['f'],'=Custos!B1')
                elif ext == 'docx':
                    self.assertEqual(item['content']['text'],content['text'])
                else:
                    self.assertEqual(item['content'],content)

    def test_table_formula_text_stays_literal(self):
        from openpyxl import load_workbook
        body,_ = export_file({'kind':'table','title':'Teste','content':{'columns':['Campo'],'rows':[['=WEBSERVICE("https://example.invalid")']]}},'xlsx')
        book=load_workbook(io.BytesIO(body))
        self.assertEqual(book.active['A2'].data_type,'s')

    def test_invalid_and_hostile_archives_are_rejected(self):
        for name,data in [('file.xlsm','AAAA'),('file.xlsx','?'),('file.docx',base64.b64encode(b'not a zip').decode())]:
            with self.assertRaises(InvalidInput): import_file(name,data)
        for path, xml in [('word/vbaProject.bin',b'data'),('word/document.xml',b'<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]>')]:
            out=io.BytesIO()
            with zipfile.ZipFile(out,'w') as z:z.writestr(path,xml)
            with self.assertRaises(InvalidInput):import_file('test.docx',base64.b64encode(out.getvalue()).decode())

    def test_workbook_limits_and_multiple_sheets_persist(self):
        with tempfile.TemporaryDirectory() as temp:
            store=Store(Path(temp)/'db')
            user=store.register('test@example.org','Teste','a-long-test-password')['id']
            ws=Workspace(store)
            content={'cells':[['1']], 'sheets':[{'name':'Primeira','cells':[['1']]},{'name':'Segunda','cells':[['2']]}], 'activeSheet':0}
            item=ws.create(user,{'kind':'sheet','title':'Abas','content':content})
            self.assertEqual(ws.get(user,item['id'])['content'],content)
            content['sheets'][1]['name']='PRIMEIRA'
            with self.assertRaises(InvalidInput):validate_content('sheet',content)

    def test_xlsx_exports_numbers_and_translates_localized_functions(self):
        from openpyxl import load_workbook
        body,_=export_file({'kind':'sheet','title':'Valores','content':{'cells':[['10','20','=SOMA(A1:B1)','=SE(1;"SOMA;";"n")']]}},'xlsx')
        book=load_workbook(io.BytesIO(body))
        self.assertEqual(book.active['A1'].value,10)
        self.assertEqual(book.active['C1'].value,'=SUM(A1:B1)')
        self.assertEqual(book.active['D1'].value,'=IF(1,"SOMA;","n")')

    def test_ics_roundtrip_and_bounded_recurrence(self):
        from server.calendar_io import import_calendar
        content={'events':[{'title':'Reunião','start':'2026-11-10T13:00:00+00:00','end':'2026-11-10T14:00:00+00:00','location':'Sala A','description':''}]}
        body,mime=export_file({'id':'example','kind':'calendar','title':'Agenda','content':content},'ics')
        item,warnings=import_file('Agenda.ics',base64.b64encode(body).decode())
        self.assertEqual(item['content'],content)
        recurring=body.replace(b'END:VEVENT',b'RRULE:FREQ=WEEKLY;COUNT=3\r\nEND:VEVENT')
        self.assertEqual(len(import_calendar(recurring)['events']),3)
        with self.assertRaises(InvalidInput):import_calendar(recurring.replace(b';COUNT=3',b''))
