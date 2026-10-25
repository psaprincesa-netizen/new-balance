import base64,copy,io,unittest
from PIL import Image
from docx import Document
from server.schema import validate_content
from server.store import InvalidInput
from server.office import export_file,import_file

class DocumentIOTest(unittest.TestCase):
    def test_real_document_tables_links_lists_and_images(self):
        out=io.BytesIO();Image.new('RGB',(160,90),'green').save(out,format='PNG')
        c={'text':'# Plano\n- Tarefa\n[Site](https://example.org)\n| Nome | Valor |\n| --- | --- |\n| Ana | 10 |\n![Foto](drap-image:photo)','images':[{'id':'photo','alt':'Foto','src':'data:image/png;base64,'+base64.b64encode(out.getvalue()).decode(),'width':60}]}
        validate_content('document',c);data,_=export_file({'kind':'document','title':'Plano','content':c},'docx');doc=Document(io.BytesIO(data))
        self.assertEqual(len(doc.tables),1);self.assertEqual(doc.tables[0].cell(1,0).text,'Ana');self.assertEqual(len(doc.inline_shapes),1);self.assertEqual(doc.paragraphs[1].style.name,'List Bullet')
        restored,_=import_file('Plano.docx',base64.b64encode(data).decode());validate_content('document',restored['content']);self.assertEqual(len(restored['content']['images']),1)
        for value in ['[Site](https://example.org)','| Nome | Valor |','- Tarefa','drap-image:docx-0']:self.assertIn(value,restored['content']['text'])
        c['images'][0]['src']='data:image/svg+xml;base64,AAAA'
        with self.assertRaises(InvalidInput):validate_content('document',c)
    def test_calendar_compares_instants_across_timezones(self):
        valid={'events':[{'title':'Reunião','start':'2026-10-25T10:00:00-03:00','end':'2026-10-25T12:30:00-01:00'}]}
        validate_content('calendar',valid)
        valid['events'][0]['end']='2026-10-25T12:30:00Z'
        with self.assertRaises(InvalidInput):validate_content('calendar',valid)

    def test_nested_marks_escaped_literals_and_table_breaks(self):
        c={'text':r'***Destaque*** e 2 \* 3'+'\n| Campo |\n| --- |\n| A<br>B |'}
        data,_=export_file({'kind':'document','title':'Texto','content':c},'docx');doc=Document(io.BytesIO(data))
        self.assertTrue(any(r.bold and r.italic and r.text=='Destaque' for r in doc.paragraphs[0].runs))
        self.assertIn('2 * 3',doc.paragraphs[0].text);self.assertEqual(doc.tables[0].cell(1,0).text,'A\nB')
        restored,_=import_file('Texto.docx',base64.b64encode(data).decode());self.assertIn('***Destaque***',restored['content']['text'])
