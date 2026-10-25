import base64,copy,io,unittest
from PIL import Image
from pptx import Presentation
from server.office import export_file,import_file
from server.schema import validate_content
from server.store import InvalidInput

def content():
    out=io.BytesIO();Image.new('RGB',(160,90),'green').save(out,format='PNG')
    return {'slides':[{'title':'Plano','body':'Resumo','notes':'Notas privadas','objects':[
        {'id':'img','type':'image','src':'data:image/png;base64,'+base64.b64encode(out.getvalue()).decode(),'alt':'Paisagem','x':55,'y':30,'w':40,'h':40},
        {'id':'text','type':'text','text':'Meta 2027','color':'#123456','fill':'#ffffff','fontSize':32,'bold':True,'align':'center','x':5,'y':70,'w':40,'h':20},
        {'id':'shape','type':'ellipse','text':'OK','color':'#ffffff','fill':'#4b5320','x':60,'y':75,'w':20,'h':20},
    ]}]}

class SlideTest(unittest.TestCase):
    def test_images_shapes_text_roundtrip(self):
        c=content();validate_content('slides',c)
        data,_=export_file({'title':'Plano','kind':'slides','content':c},'pptx');p=Presentation(io.BytesIO(data));slide=p.slides[0]
        self.assertAlmostEqual(p.slide_width/p.slide_height,16/9,places=5)
        picture=next(s for s in slide.shapes if s.name=='DRAP image');self.assertEqual(picture.image.size,(160,90))
        self.assertAlmostEqual(picture.left/p.slide_width,.55,places=4)
        restored,_=import_file('slides.pptx',base64.b64encode(data).decode());validate_content('slides',restored['content'])
        result=restored['content']['slides'][0];self.assertEqual(result['body'],'Resumo');self.assertEqual(result['notes'],'Notas privadas')
        objects=result['objects'];self.assertEqual([o['type'] for o in objects],['image','text','ellipse'])
        self.assertEqual(objects[1]['text'],'Meta 2027');self.assertEqual(objects[1]['fontSize'],32);self.assertTrue(objects[1]['bold']);self.assertEqual(objects[1]['color'].lower(),'#123456');self.assertEqual(objects[2]['fill'].lower(),'#4b5320')
    def test_invalid_images_and_geometry_rejected(self):
        c=content()
        for key,value in [('src','javascript:alert(1)'),('src','data:image/png;base64,PHN2Zz4='),('w',0),('x',95),('h',float('nan')),('type','script')]:
            bad=copy.deepcopy(c);bad['slides'][0]['objects'][0][key]=value
            with self.subTest(key=key),self.assertRaises(InvalidInput):validate_content('slides',bad)
        for key,value in [('fontSize',1000),('fill','url(https://example.org)'),('bold','yes')]:
            bad=copy.deepcopy(c);bad['slides'][0]['objects'][1][key]=value
            with self.assertRaises(InvalidInput):validate_content('slides',bad)
    def test_object_limits_and_legacy(self):
        c=content();c['slides'][0]['objects']*=11
        with self.assertRaises(InvalidInput):validate_content('slides',c)
        del c['slides'][0]['objects'];validate_content('slides',c)
