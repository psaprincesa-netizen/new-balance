"""Acceptance of real file exchange, drafts, collaboration and public collection."""
import os
import io
import tempfile
import threading
import zipfile
from pathlib import Path
from http.server import ThreadingHTTPServer
from playwright.sync_api import sync_playwright,expect
from server.store import Store
from server.app import handler_for


def main():
    with tempfile.TemporaryDirectory() as temp:
        server=ThreadingHTTPServer(('127.0.0.1',0),handler_for(Store(Path(temp)/'test.sqlite')))
        thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        try:
            with sync_playwright() as p:
                browser=p.chromium.launch(headless=True,executable_path=os.environ.get('CHROMIUM_PATH','/usr/bin/chromium'))
                page=browser.new_page(viewport={'width':1440,'height':1000});page.set_default_timeout(15000)
                errors=[];page.on('pageerror',lambda e:errors.append(str(e)));page.on('dialog',lambda d:d.accept())
                base=f'http://127.0.0.1:{server.server_port}'
                page.goto(base);expect(page.locator('.brand')).to_contain_text('DRAP office')
                page.get_by_role('button',name='Criar minha conta').click()
                page.get_by_label('Nome',exact=True).fill('Ana Teste')
                page.get_by_label('E-mail',exact=True).fill('features@example.org')
                page.get_by_label('Senha',exact=True).fill('uma-senha-para-testes-123')
                page.get_by_role('button',name='Criar conta',exact=False).click()
                expect(page.get_by_role('heading',name='Olá, Ana. Vamos criar?')).to_be_visible()
                def create(kind,title):
                    page.get_by_role('button',name='Todos os arquivos',exact=False).click()
                    page.get_by_role('button',name='＋ Criar novo').click();page.locator('#create-form select').select_option(kind)
                    page.get_by_label('Nome do arquivo',exact=True).last.fill(title);page.get_by_role('button',name='Criar arquivo',exact=True).click()
                    expect(page.locator('#title')).to_have_value(title)
                def save():
                    page.get_by_role('button',name='Salvar',exact=True).click();expect(page.locator('#save-state')).to_contain_text('Versão')
                create('document','Documento avançado')
                page.get_by_label('Conteúdo do documento').fill('# Plano\n\n**Texto** de teste.');save()
                page.get_by_role('button',name='Comentários',exact=True).click();page.get_by_label('Seu comentário').fill('Revisar a entrega')
                page.get_by_role('button',name='Enviar comentário',exact=True).click();expect(page.locator('dialog')).to_contain_text('Revisar a entrega')
                page.get_by_role('button',name='Fechar',exact=True).click()
                page.get_by_label('Conteúdo do documento').fill('Rascunho recuperável')
                for _ in range(30):
                    if page.evaluate("Object.keys(localStorage).some(k=>k.startsWith('drap:draft:'))"):break
                    page.wait_for_timeout(100)
                else:raise AssertionError('Rascunho não foi persistido')
                page.reload();page.get_by_role('button',name='Documento avançado',exact=False).click()
                expect(page.get_by_role('heading',name='Recuperar rascunho?')).to_be_visible();page.get_by_role('button',name='Confirmar',exact=True).click()
                expect(page.get_by_label('Conteúdo do documento')).to_have_value('Rascunho recuperável');save()
                with page.expect_download() as output:page.get_by_role('button',name='Baixar DOCX').click()
                doc=output.value;doc.save_as(str(Path(temp)/'documento.docx'))
                with zipfile.ZipFile(doc.path()) as z:assert 'word/document.xml' in z.namelist()
                with page.expect_file_chooser() as chooser:page.get_by_role('button',name='Importar Office / Agenda').click()
                chooser.value.set_files(str(Path(temp)/'documento.docx'))
                expect(page.get_by_role('heading',name='Arquivo importado')).to_be_visible();page.get_by_role('button',name='Continuar edição').click()
                expect(page.get_by_label('Conteúdo do documento')).to_have_value('Rascunho recuperável')
                print('PASS: marca, comentários, rascunho recuperado e DOCX real',flush=True)
                create('sheet','Pasta com abas')
                expect(page.locator('#spreadsheet')).to_have_attribute('data-ready','true')
                page.locator('#sheet-address').fill('B1');page.locator('#sheet-value').fill('25');page.locator('#sheet-apply').click()
                page.locator('#spreadsheet .univerjs-icon-increase-icon:visible').first.click()
                page.locator('#sheet-address').fill('A1');page.locator('#sheet-value').fill("='Planilha 1'!B1*2");page.locator('#sheet-apply').click();save()
                with page.expect_download() as output:page.get_by_role('button',name='Baixar XLSX').click()
                from openpyxl import load_workbook
                book=load_workbook(io.BytesIO(Path(output.value.path()).read_bytes()));assert len(book.sheetnames)==2;assert book.worksheets[1]['A1'].value=="='Planilha 1'!B1*2"
                print('PASS: múltiplas abas, referência entre abas e XLSX real',flush=True)
                create('form','Coleta pública')
                page.get_by_label('Pergunta 1',exact=True).fill('Contato')
                page.get_by_label('Tipo da pergunta 1',exact=True).select_option('email');save()
                page.get_by_role('button',name='Link de coleta').click();page.get_by_role('button',name='Ativar link público').click()
                expect(page.get_by_label('Link público')).to_be_visible();url=page.get_by_label('Link público').input_value()
                public_context=browser.new_context();public=public_context.new_page();public.goto(url)
                public.get_by_label('Contato',exact=False).fill('contato@example.org');public.get_by_role('button',name='Enviar resposta').click()
                expect(public.get_by_role('heading',name='Resposta enviada')).to_be_visible()
                page.get_by_role('button',name='Desativar link').click();expect(page.locator('dialog')).to_contain_text('fechado')
                token=url.split('?form=')[1];assert public.request.get(base+'/api/forms/'+token).status==404
                public_context.close();page.get_by_role('button',name='Fechar',exact=True).click()
                page.reload();page.get_by_role('button',name='Coleta pública',exact=False).click()
                expect(page.locator('#editor-body')).to_contain_text('contato@example.org')
                print('PASS: formulário tipado, resposta sem login, persistência e revogação',flush=True)
                create('diagram','Fluxo livre')
                node=page.get_by_role('button',name='Mover etapa Início');node.focus();node.press('ArrowRight');save()
                with page.expect_download() as output:page.get_by_role('button',name='Baixar SVG').click()
                assert 'translate(45,35)' in Path(output.value.path()).read_text()
                page.set_viewport_size({'width':390,'height':844});assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
                assert not errors,errors
                browser.close();print('PASS: canvas, teclado, SVG, móvel e ausência de erros JavaScript',flush=True)
        finally:
            server.shutdown();server.server_close();thread.join()

if __name__=='__main__':main()
