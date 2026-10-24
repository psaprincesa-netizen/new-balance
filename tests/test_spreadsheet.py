"""Acceptance contracts for native spreadsheet persistence and XLSX fidelity."""
import base64,copy,io,json,unittest
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from server.office import import_file,export_file
from server.schema import validate_content
from server.store import InvalidInput

class SpreadsheetTest(unittest.TestCase):
    def fixture(self):
        book=Workbook();s=book.active;s.title='Orçamento'
        s['A1']='Produto';s['B1']='Valor';s['A2']='Serviço';s['B2']=125.5;s['C2']='=B2*2';s['A4']='Título';s.merge_cells('A4:C4')
        s['B2'].number_format='"R$" #,##0.00';s['B2'].font=Font(bold=True,color='123456');s['B2'].fill=PatternFill('solid',fgColor='AACCAA');s['B2'].alignment=Alignment(horizontal='right',wrap_text=True)
        s.freeze_panes='B2';s.column_dimensions['B'].width=25;s.row_dimensions[2].height=30
        book.create_sheet('Resumo')['A1']="='Orçamento'!C2"
        out=io.BytesIO();book.save(out)
        item,_=import_file('Teste.xlsx',base64.b64encode(out.getvalue()).decode());return item

    def test_styles_formulas_merges_dimensions_and_freeze_roundtrip(self):
        item=self.fixture();validate_content('sheet',item['content'])
        data,_=export_file(item,'xlsx');book=load_workbook(io.BytesIO(data));s=book['Orçamento']
        self.assertEqual(book.sheetnames,['Orçamento','Resumo']);self.assertEqual(s['C2'].value,'=B2*2');self.assertEqual(s['B2'].value,125.5)
        self.assertTrue(s['B2'].font.bold);self.assertEqual(s['B2'].fill.fgColor.rgb[-6:],'AACCAA');self.assertEqual(s['B2'].number_format,'"R$" #,##0.00')
        self.assertIn('A4:C4',s.merged_cells);self.assertEqual(s.freeze_panes,'B2');self.assertEqual(s.column_dimensions['B'].width,25);self.assertEqual(s.row_dimensions[2].height,30)
        self.assertTrue(s['B2'].alignment.wrap_text)

    def test_native_resources_export_as_excel_rules(self):
        item=self.fixture();b=item['content']['workbook'];r={'startRow':1,'endRow':10,'startColumn':1,'endColumn':1}
        resources={
            'SHEET_FILTER_PLUGIN':{'sheet-0':{'ref':{**r,'startRow':0,'startColumn':0},'filterColumns':[{'colId':0,'filters':{'filters':['Serviço']}}]}},
            'SHEET_DATA_VALIDATION_PLUGIN':{'sheet-0':[{'uid':'v','type':'decimal','operator':'greaterThan','formula1':'0','ranges':[r],'errorStyle':1,'showErrorMessage':True}]},
            'SHEET_CONDITIONAL_FORMATTING_PLUGIN':{'sheet-0':[{'cfId':'c','ranges':[r],'rule':{'type':'highlightCell','subType':'number','operator':'greaterThan','value':100,'style':{'bg':{'rgb':'#FF0000'}}}}]},
        }
        b['resources']=[{'name':k,'data':json.dumps(v)} for k,v in resources.items()];validate_content('sheet',item['content'])
        data,_=export_file(item,'xlsx');s=load_workbook(io.BytesIO(data)).worksheets[0]
        self.assertEqual(s.auto_filter.ref,'A1:B11');self.assertEqual(s.data_validations.dataValidation[0].formula1,'0');self.assertEqual(s.data_validations.dataValidation[0].sqref.__str__(),'B2:B11')
        rules=list(s.conditional_formatting);self.assertEqual(len(rules),1);self.assertEqual(s.conditional_formatting[rules[0]][0].operator,'greaterThan')
        restored,_=import_file('Regras.xlsx',base64.b64encode(data).decode())
        r={entry['name']:json.loads(entry['data']) for entry in restored['content']['workbook']['resources']}
        self.assertEqual(r['SHEET_DATA_VALIDATION_PLUGIN']['sheet-0'][0]['formula1'],'0')
        self.assertEqual(r['SHEET_CONDITIONAL_FORMATTING_PLUGIN']['sheet-0'][0]['rule']['value'],100)
        validate_content('sheet',restored['content'])

    def test_sparse_bounds_and_duplicate_names_are_rejected(self):
        content=self.fixture()['content']
        for modify in [lambda b:b['sheets']['sheet-0']['cellData'].update({'100000':{'0':{'v':'bad'}}}),lambda b:b['sheets']['sheet-1'].update(name='ORÇAMENTO'),lambda b:b.update(sheetOrder=['sheet-0','missing']),lambda b:b['sheets']['sheet-0'].update(rowCount=100001)]:
            candidate=copy.deepcopy(content);modify(candidate['workbook'])
            with self.assertRaises(InvalidInput):validate_content('sheet',candidate)

    def test_1904_dates_do_not_shift_on_roundtrip(self):
        from datetime import datetime
        from openpyxl.utils.datetime import CALENDAR_MAC_1904
        book=Workbook();book.epoch=CALENDAR_MAC_1904;book.active['A1']=datetime(2026,10,24)
        data=io.BytesIO();book.save(data)
        item,_=import_file('Data.xlsx',base64.b64encode(data.getvalue()).decode())
        output,_=export_file(item,'xlsx');result=load_workbook(io.BytesIO(output))
        self.assertEqual(result.epoch,CALENDAR_MAC_1904);self.assertEqual(result.active['A1'].value,datetime(2026,10,24))

    def test_charts_export_real_excel_series(self):
        item=self.fixture();item['content']['charts']=[{'id':'c','type':'bar','title':'Custos','sheetId':'sheet-0','range':'A1:B2'}]
        validate_content('sheet',item['content']);data,_=export_file(item,'xlsx')
        result=load_workbook(io.BytesIO(data));chart=result.worksheets[0]._charts[0]
        self.assertEqual(chart.series[0].val.numRef.f,"'Orçamento'!$B$2")
        item['content']['charts'][0]['range']='A1:XFD1000000'
        with self.assertRaises(InvalidInput):validate_content('sheet',item['content'])
