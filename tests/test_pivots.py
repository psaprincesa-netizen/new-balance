import io,json,copy,unittest
from pathlib import Path
from openpyxl import load_workbook
from server.pivots import calculate_pivot,valid_pivots,parse_range
from server.schema import validate_content
from server.office import export_file
from server.store import InvalidInput
CASES=json.loads(Path(__file__).with_name('pivot-cases.json').read_text())

def fixture():
    source=CASES[0]['source'];cells={str(r):{str(c):{'v':v} for c,v in enumerate(row) if v is not None} for r,row in enumerate(source)}
    p={**CASES[0]['config'],'id':'pivot-test','title':'Resumo de receitas','aggregate':'sum','sheetId':'s','range':'A1:C7'}
    return {'workbook':{'id':'book','sheetOrder':['s'],'styles':{},'sheets':{'s':{'id':'s','name':'Dados','rowCount':100,'columnCount':26,'cellData':cells}}},'pivots':[p]}

class PivotTest(unittest.TestCase):
    def test_shared_aggregation_cases(self):
        for c in CASES:
            with self.subTest(c['name']):self.assertEqual(calculate_pivot(c['source'],c['config']),c['expected'])
    def test_validation_rejects_bad_definitions(self):
        c=fixture();validate_content('sheet',c)
        for key,value in [('aggregate',[]),('rowField',True),('valueField',99),('range','A1:ZZ99999'),('sort','malicious'),('title','')]:
            bad=copy.deepcopy(c);bad['pivots'][0][key]=value
            with self.subTest(key=key),self.assertRaises(InvalidInput):validate_content('sheet',bad)
        self.assertFalse(valid_pivots(c['pivots']*11))
    def test_xlsx_exports_current_results_and_literal_labels(self):
        from server.spreadsheet_io import export_workbook
        output=io.BytesIO();export_workbook(fixture()['workbook'],output)
        self.assertEqual(load_workbook(io.BytesIO(output.getvalue())).sheetnames,['Dados'])
        c=fixture();c['workbook']['sheets']['s']['cellData']['1']['2']['p']=None;data,_=export_file({'kind':'sheet','title':'Dados','content':c},'xlsx');book=load_workbook(io.BytesIO(data))
        self.assertEqual(list(book['Resumo 1'].values),[tuple(row) for row in CASES[0]['expected']])
        c['workbook']['sheets']['s']['cellData']['1']['0']['v']='=1+1'
        data,_=export_file({'kind':'sheet','title':'Dados','content':c},'xlsx');book=load_workbook(io.BytesIO(data));self.assertEqual(book['Resumo 1']['A2'].data_type,'s')
    def test_deleted_source_and_formula_errors_fail_explicitly(self):
        c=fixture();c['pivots'][0]['sheetId']='missing'
        validate_content('sheet',c) # keep broken definitions recoverable after sheet deletion
        with self.assertRaises(InvalidInput):export_file({'kind':'sheet','title':'Dados','content':c},'xlsx')
        c=fixture();c['workbook']['sheets']['s']['cellData']['1']['2']={'f':'=1/0','v':'#DIV/0!'}
        with self.assertRaises(InvalidInput):export_file({'kind':'sheet','title':'Dados','content':c},'xlsx')
    def test_overflow_and_duplicate_headers(self):
        p={'rowField':0,'valueField':1,'aggregate':'sum'}
        for rows in [[['A','B'],['x',1e308],['x',1e308]],[['A','A'],['x',1]]]:
            with self.assertRaises(InvalidInput):calculate_pivot(rows,p)
