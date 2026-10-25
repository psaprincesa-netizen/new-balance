"""Create, manipulate, save, present and exchange slide images and objects in a browser."""
import io,os,tempfile,threading,uuid,base64
from pathlib import Path
from http.server import ThreadingHTTPServer
from playwright.sync_api import sync_playwright,expect
from PIL import Image
from pptx import Presentation
from server.app import handler_for
from server.store import Store

def exercise(page,base,title='Apresentação visual'):
    errors=[];page.on('pageerror',lambda e:(errors.append(str(e)),print('Browser error:',e,flush=True)))
    page.goto(base);page.get_by_role('button',name='Criar minha conta').click()
    page.get_by_label('Nome',exact=True).fill('Ana Teste');page.get_by_label('E-mail',exact=True).fill(uuid.uuid4().hex+'@example.org');page.get_by_label('Senha',exact=True).fill(uuid.uuid4().hex+'-Teste12')
    page.get_by_role('button',name='Criar conta',exact=False).click();expect(page.get_by_role('heading',name='Olá, Ana. Vamos criar?')).to_be_visible()
    page.get_by_role('button',name='＋ Criar novo').click();page.locator('#create-form select').select_option('slides');page.get_by_label('Nome do arquivo',exact=True).last.fill(title);page.get_by_role('button',name='Criar arquivo',exact=True).click()
    page.locator('#slide-title').fill('Resultados');page.locator('#slide-body').fill('Resumo do trimestre');page.locator('#slide-notes').fill('Notas para apresentar')
    out=io.BytesIO();Image.new('RGB',(320,180),'#4b5320').save(out,format='PNG')
    with page.expect_file_chooser() as chooser:page.get_by_role('button',name='Inserir imagem',exact=True).click()
    chooser.value.set_files({'name':'foto.png','mimeType':'image/png','buffer':out.getvalue()})
    expect(page.locator('#slide-preview img')).to_be_visible();page.get_by_label('Descrição da imagem',exact=True).fill('Imagem da apresentação');page.get_by_label('Descrição da imagem',exact=True).press('Tab')
    image=page.locator('#slide-preview .image');image.scroll_into_view_if_needed();box=image.bounding_box();page.mouse.move(box['x']+30,box['y']+30);page.mouse.down();page.mouse.move(box['x']+5,box['y']+15,steps=5);page.mouse.up()
    handle=page.get_by_role('button',name='Redimensionar objeto',exact=True);handle.scroll_into_view_if_needed();box=handle.bounding_box();page.mouse.move(box['x']+10,box['y']+10);page.mouse.down();page.mouse.move(box['x']+30,box['y']+30,steps=5);page.mouse.up()
    page.get_by_role('button',name='Caixa de texto',exact=True).click();page.get_by_label('Texto do objeto',exact=True).fill('Meta 2027');page.get_by_label('Texto do objeto',exact=True).press('Tab');page.get_by_label('Tamanho do texto',exact=True).fill('32');page.get_by_label('Tamanho do texto',exact=True).press('Tab');page.locator('#object-bold').check();page.get_by_role('button',name='Retângulo',exact=True).click()
    page.get_by_role('button',name='Excluir objeto',exact=True).click();page.get_by_role('button',name='Desfazer alteração',exact=True).click();expect(page.locator('#slide-preview .rect')).to_have_count(1)
    options=page.locator('#slide-object-list option').all();page.locator('#slide-object-list').select_option(options[2].get_attribute('value'));page.get_by_role('button',name='Trazer à frente',exact=True).click();page.get_by_label('Posição horizontal do objeto',exact=True).fill('6');page.get_by_label('Posição horizontal do objeto',exact=True).press('Tab');page.get_by_label('Posição vertical do objeto',exact=True).fill('70');page.get_by_label('Posição vertical do objeto',exact=True).press('Tab')
    page.get_by_role('button',name='＋ Slide',exact=True).click();page.locator('#slide-title').fill('Conclusão');page.get_by_role('button',name='Mover antes',exact=True).click();expect(page.locator('[data-slide="0"]')).to_contain_text('Conclusão');page.get_by_role('button',name='Mover depois',exact=True).click()
    page.get_by_role('button',name='Duplicar slide',exact=True).click();expect(page.locator('[data-slide]')).to_have_count(3);page.get_by_role('button',name='Excluir slide',exact=True).click();expect(page.locator('[data-slide]')).to_have_count(2)
    page.get_by_role('button',name='Salvar',exact=True).click();expect(page.locator('#save-state')).to_contain_text('Versão');expect(page.locator('#slide-title')).to_have_value('Conclusão')
    page.reload();page.get_by_role('button',name=title,exact=False).click();expect(page.locator('#slide-preview img')).to_be_visible()
    page.get_by_role('button',name='▷ Apresentar',exact=True).click();expect(page.locator('.present img')).to_be_visible();page.keyboard.press('ArrowRight');expect(page.locator('.present h1')).to_have_text('Conclusão');page.keyboard.press('Escape')
    page.evaluate('() => {window.print=()=>{window.printCheck={slides:document.querySelectorAll(".print-slide-deck .slide-stage").length,images:document.querySelectorAll(".print-slide-deck img").length,text:document.querySelector(".print-slide-deck").textContent};};}')
    page.get_by_role('button',name='Imprimir / PDF',exact=True).click();printed=page.evaluate('() => window.printCheck');assert printed['slides']==2 and printed['images']==1 and 'Notas para apresentar' not in printed['text']
    with page.expect_download() as dl:page.get_by_role('button',name='Baixar PPTX',exact=True).click()
    data=Path(dl.value.path()).read_bytes();p=Presentation(io.BytesIO(data));assert len(p.slides)==2;assert any(s.name=='DRAP image' for s in p.slides[0].shapes)
    item=next(x for x in page.request.get(base+'/api/items').json()['items'] if x['title']==title);objects=item['content']['slides'][0]['objects'];assert len(objects)==3;assert objects[0]['x']<54 and objects[0]['w']>40,{k:objects[0][k] for k in ['x','y','w','h']};text=next(o for o in objects if o['type']=='text');assert text['bold'] and text['fontSize']==32 and text['x']==6
    with page.expect_file_chooser() as chooser:page.get_by_role('button',name='Importar Office / Agenda',exact=True).click()
    chooser.value.set_files({'name':'Reimportado.pptx','mimeType':'application/vnd.openxmlformats-officedocument.presentationml.presentation','buffer':data});page.get_by_role('button',name='Continuar edição',exact=True).click();expect(page.locator('#slide-preview img')).to_be_visible();expect(page.locator('#slide-preview')).to_contain_text('Meta 2027')
    imported=next(x for x in page.request.get(base+'/api/items').json()['items'] if x['title']=='Reimportado')
    assert page.request.post(base+'/api/items/'+imported['id'],headers={'Origin':base},data={'title':imported['title'],'content':imported['content'],'version':imported['version'],'deleted':True}).ok
    page.get_by_role('button',name='Todos os arquivos',exact=False).click();page.get_by_role('button',name=title,exact=False).click()
    expect(page.locator('#slide-preview img')).to_be_visible()
    page.screenshot(path='/tmp/drap-slides-verified.png',full_page=True)
    ctx=page.context.browser.new_context();reader=ctx.new_page();credentials={'email':uuid.uuid4().hex+'@example.org','password':uuid.uuid4().hex+'-Reader12'};headers={'Origin':base}
    assert reader.request.post(base+'/api/register',headers=headers,data={**credentials,'name':'Leitor'}).ok
    assert reader.request.post(base+'/api/login',headers=headers,data=credentials).ok
    assert page.request.post(base+'/api/items/'+item['id']+'/share',headers=headers,data={'email':credentials['email'],'role':'read'}).ok
    reader.goto(base);reader.get_by_role('button',name=title,exact=False).click();expect(reader.locator('#slide-preview img')).to_be_visible();expect(reader.locator('#slide-image')).to_be_disabled();reader.locator('[data-slide="1"]').click();expect(reader.locator('#slide-image')).to_be_disabled();expect(reader.locator('#slide-title')).to_have_value('Conclusão');reader.get_by_role('button',name='▷ Apresentar',exact=True).click();expect(reader.locator('.present img')).to_be_visible();reader.keyboard.press('Escape')
    denied=reader.request.post(base+'/api/items/'+item['id'],headers=headers,data={'title':title,'content':item['content'],'version':item['version'],'deleted':False});assert denied.status==403
    assert page.request.post(base+'/api/items/'+item['id']+'/share',headers=headers,data={'email':credentials['email'],'role':'remove'}).ok
    reader.request.post(base+'/api/logout',headers=headers,data={});ctx.close()
    page.set_viewport_size({'width':390,'height':844});expect(page.get_by_role('button',name='Inserir imagem',exact=True)).to_be_visible();expect(page.locator('#slide-preview img')).to_be_visible()
    assert not errors,errors
    print('PASS: slides, imagem, arrastar/redimensionar, texto, forma, desfazer, reordenar, salvar, apresentar, PPTX e reimportação, celular',flush=True)
    return item

def main():
    expect.set_options(timeout=20000)
    with tempfile.TemporaryDirectory() as tmp:
        server=ThreadingHTTPServer(('127.0.0.1',0),handler_for(Store(Path(tmp)/'test.sqlite')));threading.Thread(target=server.serve_forever,daemon=True).start()
        try:
            with sync_playwright() as p:
                browser=p.chromium.launch(headless=True,executable_path=os.environ.get('CHROMIUM_PATH','/usr/bin/chromium'));page=browser.new_page(viewport={'width':1600,'height':1100});page.set_default_timeout(20000);exercise(page,f'http://127.0.0.1:{server.server_port}');browser.close()
        finally:server.shutdown();server.server_close()
if __name__=='__main__':main()
