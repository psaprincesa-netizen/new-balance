import os
"""Real UI acceptance: native editing, formatting, rules, persistence and XLSX."""
import io,json,tempfile,threading,uuid
from pathlib import Path
from http.server import ThreadingHTTPServer
from playwright.sync_api import sync_playwright,expect
from server.app import handler_for
from server.store import Store

def exercise(page,base,title='Planilha de aceitação'):
    errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    page.goto(base);page.get_by_role('button',name='Criar minha conta').click()
    page.get_by_label('Nome',exact=True).fill('Ana Teste');page.get_by_label('E-mail',exact=True).fill(uuid.uuid4().hex+'@example.org');page.get_by_label('Senha',exact=True).fill(uuid.uuid4().hex+'-Teste12')
    page.get_by_role('button',name='Criar conta',exact=False).click();expect(page.get_by_role('heading',name='Olá, Ana. Vamos criar?')).to_be_visible()
    page.get_by_role('button',name='＋ Criar novo').click();page.locator('#create-form select').select_option('sheet');page.get_by_label('Nome do arquivo',exact=True).last.fill(title);page.get_by_role('button',name='Criar arquivo',exact=True).click()
    def ready():expect(page.locator('#spreadsheet')).to_have_attribute('data-ready','true')
    def select(address):
        page.locator('#sheet-address').fill(address);page.locator('#sheet-select').click()
    def setcell(address,value):
        page.locator('#sheet-address').fill(address);page.locator('#sheet-value').fill(value);page.locator('#sheet-apply').click()
    def save():
        page.get_by_role('button',name='Salvar',exact=True).click();expect(page.locator('#save-state')).to_contain_text('Versão');ready()
    ready()
    # Type directly into the canvas with the keyboard, without calling a facade.
    canvas=None
    for _ in range(100):
        canvas=next((c for c in page.locator('#spreadsheet canvas').all() if c.bounding_box() and c.bounding_box()['height']>300),None)
        if canvas:break
        page.wait_for_timeout(100)
    assert canvas is not None, 'A grade não abriu'
    canvas.click(position={'x':70,'y':35});page.keyboard.type('Produto');page.keyboard.press('Enter')
    for addr,val in [('B1','Valor'),('A2','B'),('B2','20'),('A3','A'),('B3','10'),('C2','=SOMA(B2:B3)')]:setcell(addr,val)
    select('A1:B1');page.locator('button[aria-label="Negrito"]:visible').click()
    page.locator('button[aria-label="Desfazer"]:visible').click();page.locator('button[aria-label="Refazer"]:visible').click()
    select('D4:E4');page.locator('[aria-label="Mesclar células"]:visible').click();page.get_by_text('Mesclar tudo',exact=True).click()
    page.locator('#spreadsheet').get_by_text('Dados',exact=True).click()
    select('A2:B3');page.locator('#spreadsheet .univerjs-icon-ascending-icon:visible').click();page.get_by_text('Crescente',exact=True).click()
    select('A1:B3');page.locator('#spreadsheet .univerjs-icon-filter-icon:visible').click();page.keyboard.press('Escape')
    select('B2:B3');page.locator('[aria-label="Validação de dados"]:visible').click();page.get_by_text('Adicionar regra',exact=True).click();page.get_by_text('Concluir',exact=True).click();page.get_by_role('button',name='Fechar o painel lateral').click()
    page.locator('[aria-label="Formatação Condicional"]:visible').click();page.get_by_text('Criar Formatação Condicional',exact=True).click()
    page.get_by_text('Texto',exact=True).click();page.get_by_text('Número',exact=True).click()
    panel=page.locator('[aria-label="Painel lateral"]')
    panel.locator('input').nth(1).fill('15');page.get_by_text('Aplicar',exact=True).click()
    if page.get_by_role('button',name='Fechar o painel lateral').is_visible():page.get_by_role('button',name='Fechar o painel lateral').click()
    page.set_viewport_size({'width':1920,'height':1100});page.get_by_role('tab',name='Início',exact=True).click();select('B2:B3')
    page.locator('[data-u-command="sheet.command.numfmt.set.currency"]:visible').click();page.set_viewport_size({'width':1600,'height':1100})
    select('A1:B3');page.get_by_role('button',name='Criar gráfico',exact=True).click()
    page.get_by_label('Título do gráfico',exact=True).fill('Valores por produto');page.get_by_label('Intervalo de dados',exact=True).fill('A1:B3');page.get_by_role('button',name='Adicionar gráfico',exact=True).click()
    expect(page.locator('#sheet-charts svg')).to_be_visible()
    setcell('B3','25');expect(page.locator('#sheet-charts')).to_contain_text('B: 25');setcell('B3','20')
    with page.expect_download() as svg:page.get_by_role('button',name='Baixar SVG',exact=True).click()
    assert '<svg' in Path(svg.value.path()).read_text()
    save()
    item=next(x for x in page.request.get(base+'/api/items').json()['items'] if x['title']==title)
    b=item['content']['workbook'];sheet=b['sheets'][b['sheetOrder'][0]];cells=sheet['cellData']
    assert cells['0']['0']['v']=='Produto',cells
    assert b['styles'][cells['0']['0']['s']]['bl']==1,b['styles']
    assert cells['1']['2']['v']==30,cells
    assert cells['1']['0']['v']=='A' and cells['2']['0']['v']=='B',cells
    assert sheet['mergeData'],sheet['mergeData']
    resources={r['name']:json.loads(r['data'] or '{}') for r in b['resources']}
    for plugin in ['SHEET_FILTER_PLUGIN','SHEET_DATA_VALIDATION_PLUGIN','SHEET_CONDITIONAL_FORMATTING_PLUGIN']:assert resources[plugin],resources
    page.reload();page.get_by_role('button',name=title,exact=False).click();ready()
    with page.expect_download() as dl:page.get_by_role('button',name='Baixar XLSX').click()
    from openpyxl import load_workbook
    book=load_workbook(io.BytesIO(Path(dl.value.path()).read_bytes()));ws=book.worksheets[0]
    assert ws['A1'].value=='Produto' and ws['A1'].font.bold
    assert ws['C2'].value=='=SUM(B2:B3)' and 'D4:E4' in ws.merged_cells
    assert ws.auto_filter.ref=='A1:B3' and len(ws.data_validations.dataValidation)==1
    assert len(ws.conditional_formatting)==1
    assert len(ws._charts)==1 and ws._charts[0].series[0].val.numRef.f=="'Planilha 1'!$B$2:$B$3"
    expect(page.locator('#sheet-charts svg')).to_be_visible()
    with page.expect_download() as backup_download:page.get_by_role('button',name='Exportar JSON',exact=True).click()
    backup=json.loads(Path(backup_download.value.path()).read_text());backup['items'][0]['title']='Cópia JSON '+title
    with page.expect_file_chooser() as chooser:page.get_by_role('button',name='Importar JSON',exact=True).click()
    chooser.value.set_files({'name':'backup.json','mimeType':'application/json','buffer':json.dumps(backup).encode()})
    expect(page.get_by_role('status')).to_contain_text('1 arquivos importados')
    page.get_by_role('button',name='Todos os arquivos',exact=False).click();page.get_by_role('button',name='Cópia JSON '+title,exact=False).click();ready()
    expect(page.locator('#sheet-charts svg')).to_be_visible();expect(page.locator('button[aria-label="Negrito"]:visible')).to_be_enabled()
    copied=next(x for x in page.request.get(base+'/api/items').json()['items'] if x['title']=='Cópia JSON '+title)
    assert copied['content']==backup['items'][0]['content']
    assert page.request.post(base+'/api/items/'+copied['id'],headers={'Origin':base},data={'title':copied['title'],'content':copied['content'],'version':copied['version'],'deleted':True}).ok
    reader_context=page.context.browser.new_context();reader=reader_context.new_page()
    credentials={'email':uuid.uuid4().hex+'@example.org','password':uuid.uuid4().hex+'-Reader12'};headers={'Origin':base}
    assert reader.request.post(base+'/api/register',headers=headers,data={**credentials,'name':'Leitor de teste'}).ok
    assert reader.request.post(base+'/api/login',headers=headers,data=credentials).ok
    assert page.request.post(base+'/api/items/'+item['id']+'/share',headers=headers,data={'email':credentials['email'],'role':'read'}).ok
    reader.goto(base);reader.get_by_role('button',name=title,exact=False).click()
    expect(reader.locator('#spreadsheet')).to_have_attribute('data-ready','true')
    expect(reader.locator('#save')).to_be_disabled();expect(reader.locator('#sheet-apply')).to_be_disabled()
    expect(reader.locator('button[aria-label="Negrito"]:visible')).to_be_disabled()
    with reader.expect_download() as reader_backup:reader.get_by_role('button',name='Exportar JSON',exact=True).click()
    assert json.loads(Path(reader_backup.value.path()).read_text())['items'][0]['content']==item['content']
    denied=reader.request.post(base+'/api/items/'+item['id'],headers=headers,data={'title':title,'content':item['content'],'version':item['version'],'deleted':False});assert denied.status==403
    assert page.request.post(base+'/api/items/'+item['id']+'/share',headers=headers,data={'email':credentials['email'],'role':'remove'}).ok
    reader.request.post(base+'/api/logout',headers=headers,data={});reader_context.close()
    page.get_by_role('button',name='Todos os arquivos',exact=False).click();page.get_by_role('button',name=title,exact=False).click();ready()
    select('F1');page.get_by_role('button',name='Colar tabela',exact=True).click();page.get_by_label('Dados copiados do Excel ou CSV',exact=True).fill('Produto\tQuantidade\nX\t5\nY\t7');page.get_by_role('button',name='Colar dados',exact=True).click();expect(page.locator('dialog')).not_to_be_visible()
    select('G2:G3');page.get_by_role('button',name='∑ Inserir cálculo',exact=True).click();page.get_by_label('Operação',exact=True).select_option('AVERAGE');page.get_by_label('Célula do resultado',exact=True).fill('G2');page.get_by_role('button',name='Inserir fórmula',exact=True).click();expect(page.locator('#calc-error')).to_contain_text('fora da origem');page.get_by_label('Célula do resultado',exact=True).fill('H1');page.get_by_role('button',name='Inserir fórmula',exact=True).click();expect(page.locator('dialog')).not_to_be_visible()
    setcell('G3','9');select('F1');page.get_by_role('button',name='Inserir linha',exact=True).click();select('F1');page.get_by_role('button',name='Inserir coluna',exact=True).click();save()
    item=next(x for x in page.request.get(base+'/api/items').json()['items'] if x['id']==item['id'])
    b=item['content']['workbook'];cells=b['sheets'][b['sheetOrder'][0]]['cellData'];assert cells['1']['8']['v']==7,cells
    select('H3');page.get_by_role('button',name='Limpar valores',exact=True).click();save()
    item=next(x for x in page.request.get(base+'/api/items').json()['items'] if x['id']==item['id'])
    b=item['content']['workbook'];cells=b['sheets'][b['sheetOrder'][0]]['cellData'];assert cells['1']['8']['v']==9,cells
    assert not errors,errors
    page.screenshot(path='/tmp/drap-spreadsheet-verified.png',full_page=True)
    print('PASS: edição nativa, fórmula, negrito, desfazer/refazer, mesclagem, ordenação, filtro, validação, formatação condicional, gráficos vinculados/SVG, backup JSON, modo leitura, reabertura e XLSX',flush=True)
    return item

def main():
    expect.set_options(timeout=15000)
    with tempfile.TemporaryDirectory() as temp:
        server=ThreadingHTTPServer(('127.0.0.1',0),handler_for(Store(Path(temp)/'test.sqlite')));threading.Thread(target=server.serve_forever,daemon=True).start()
        try:
            with sync_playwright() as p:
                browser=p.chromium.launch(headless=True,executable_path=os.environ.get('CHROMIUM_PATH','/usr/bin/chromium'));page=browser.new_page(viewport={'width':1600,'height':1100});page.set_default_timeout(15000)
                exercise(page,f'http://127.0.0.1:{server.server_port}');browser.close()
        finally:server.shutdown();server.server_close()
if __name__=='__main__':main()
