"""Visual document editing, table commands, clipboard boundaries and persistence."""
import io,os,tempfile,threading,uuid
from pathlib import Path
from http.server import ThreadingHTTPServer
from playwright.sync_api import sync_playwright,expect
from PIL import Image
from docx import Document
from server.app import handler_for
from server.store import Store

def exercise(page,base,title='Documento visual'):
    errors=[];page.on('pageerror',lambda e:(errors.append(str(e)),print('Browser error:',e,flush=True)))
    page.goto(base);page.get_by_role('button',name='Criar minha conta').click();page.get_by_label('Nome',exact=True).fill('Ana Teste');page.get_by_label('E-mail',exact=True).fill(uuid.uuid4().hex+'@example.org');page.get_by_label('Senha',exact=True).fill(uuid.uuid4().hex+'-Teste12');page.get_by_role('button',name='Criar conta',exact=False).click();expect(page.get_by_role('heading',name='Olá, Ana. Vamos criar?')).to_be_visible()
    page.get_by_role('button',name='＋ Criar novo').click();page.locator('#create-form select').select_option('document');page.get_by_label('Nome do arquivo',exact=True).last.fill(title);page.get_by_role('button',name='Criar arquivo',exact=True).click();expect(page.locator('#preview')).to_have_attribute('data-ready','true')
    visual=page.get_by_role('textbox',name='Página do documento',exact=True);expect(visual).to_have_attribute('contenteditable','true');visual.fill('Texto importante');visual.press('Control+A');page.get_by_role('button',name='Negrito',exact=True).click();page.get_by_role('button',name='Itálico',exact=True).click();expect(visual.locator('strong em,em strong')).to_have_text('Texto importante')
    page.locator('#document-undo').click();page.locator('#document-redo').click();expect(visual.locator('strong em,em strong')).to_have_text('Texto importante')
    for search,replacement in [('importante','essencial'),('essencial','importante')]:
        page.get_by_role('button',name='Localizar e substituir',exact=True).click();page.get_by_label('Localizar',exact=True).fill(search);page.get_by_label('Substituir por',exact=True).fill(replacement);page.get_by_role('button',name='Substituir todas',exact=True).click();expect(visual.locator('strong em,em strong')).to_have_text('Texto '+replacement)
    visual.focus();visual.press('Control+End');visual.press('Enter');page.get_by_role('button',name='Tabela',exact=True).click();expect(visual.locator('table')).to_be_visible();visual.locator('th').first.click();page.get_by_role('button',name='Adicionar linha',exact=True).click();page.get_by_role('button',name='Adicionar coluna',exact=True).click();expect(visual.locator('tr')).to_have_count(3);expect(visual.locator('tr').first.locator('th,td')).to_have_count(3)
    page.get_by_role('button',name='Excluir coluna',exact=True).click();page.get_by_role('button',name='Excluir linha',exact=True).click();expect(visual.locator('tr')).to_have_count(2)
    image=io.BytesIO();Image.new('RGB',(160,90),'green').save(image,format='PNG')
    with page.expect_file_chooser() as chooser:page.get_by_role('button',name='Inserir imagem',exact=True).click()
    chooser.value.set_files({'name':'visual.png','mimeType':'image/png','buffer':image.getvalue()});expect(visual.locator('img')).to_be_visible()
    page.get_by_role('button',name='Localizar e substituir',exact=True).click();page.get_by_label('Localizar',exact=True).fill('drap-image');page.get_by_label('Substituir por',exact=True).fill('broken');page.get_by_role('button',name='Substituir todas',exact=True).click();expect(page.get_by_role('status')).to_contain_text('0 ocorrências');expect(visual.locator('img')).to_be_visible()
    visual.focus();visual.press('Control+End');visual.press('Enter')
    visual.evaluate('el => {const data=new DataTransfer();data.setData("text/plain","Texto colado <literal>");data.setData("text/html","<script>window.pasteExecuted=true</script><img src=https://example.invalid/tracker><b>Conteúdo hostil</b>");el.dispatchEvent(new ClipboardEvent("paste",{clipboardData:data,bubbles:true,cancelable:true}));}')
    expect(visual).to_contain_text('Texto colado <literal>');expect(visual.locator('img')).to_have_count(1);assert not page.evaluate('() => Boolean(window.pasteExecuted)')
    page.get_by_role('button',name='Editar Markdown',exact=True).click();expect(page.locator('#text')).to_be_visible();assert 'drap-image:' in page.locator('#text').input_value();page.get_by_role('button',name='Editar visualmente',exact=True).click();expect(visual.locator('img')).to_be_visible()
    page.get_by_role('button',name='Salvar',exact=True).click();expect(page.locator('#save-state')).to_contain_text('Versão');page.reload();page.get_by_role('button',name=title,exact=False).click();expect(page.locator('#preview')).to_have_attribute('data-ready','true');expect(visual.locator('img')).to_be_visible();expect(visual.locator('strong em,em strong')).to_have_text('Texto importante')
    expect(visual.locator('tr')).to_have_count(2);expect(page.locator('#word-count')).to_have_text('6 palavras')
    with page.expect_download() as dl:page.get_by_role('button',name='Baixar DOCX',exact=True).click()
    doc=Document(io.BytesIO(Path(dl.value.path()).read_bytes()));assert len(doc.tables)==1 and len(doc.inline_shapes)==1;assert any(run.bold and run.italic and 'Texto importante' in run.text for p in doc.paragraphs for run in p.runs)
    item=next(i for i in page.request.get(base+'/api/items').json()['items'] if i['title']==title)
    ctx=page.context.browser.new_context();reader=ctx.new_page();creds={'email':uuid.uuid4().hex+'@example.org','password':uuid.uuid4().hex+'-Reader12'};headers={'Origin':base};assert reader.request.post(base+'/api/register',headers=headers,data={**creds,'name':'Leitor'}).ok;assert reader.request.post(base+'/api/login',headers=headers,data=creds).ok;assert page.request.post(base+'/api/items/'+item['id']+'/share',headers=headers,data={'email':creds['email'],'role':'read'}).ok
    reader.goto(base);reader.get_by_role('button',name=title,exact=False).click();expect(reader.locator('#preview')).to_have_attribute('data-ready','true');expect(reader.get_by_role('textbox',name='Página do documento',exact=True)).to_have_attribute('contenteditable','false');expect(reader.get_by_role('button',name='Negrito',exact=True)).to_be_disabled();reader.get_by_role('button',name='Editar Markdown',exact=True).click();expect(reader.locator('#text')).to_be_disabled()
    assert page.request.post(base+'/api/items/'+item['id']+'/share',headers=headers,data={'email':creds['email'],'role':'remove'}).ok;reader.request.post(base+'/api/logout',headers=headers,data={});ctx.close()
    page.screenshot(path='/tmp/drap-document-verified.png',full_page=True);assert not errors,errors
    print('PASS: edição visual, negrito/itálico, desfazer/refazer, tabelas, imagem, colagem segura, Markdown, persistência, DOCX e leitura',flush=True)
    return item

def main():
    expect.set_options(timeout=20000)
    with tempfile.TemporaryDirectory() as tmp:
        server=ThreadingHTTPServer(('127.0.0.1',0),handler_for(Store(Path(tmp)/'test.sqlite')));threading.Thread(target=server.serve_forever,daemon=True).start()
        try:
            with sync_playwright() as p:
                browser=p.chromium.launch(headless=True,executable_path=os.environ.get('CHROMIUM_PATH','/usr/bin/chromium'));page=browser.new_page(viewport={'width':1600,'height':1100});page.set_default_timeout(20000);exercise(page,f'http://127.0.0.1:{server.server_port}');browser.close()
        finally:server.shutdown();server.server_close()
if __name__=='__main__':main()
